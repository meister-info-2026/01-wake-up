import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 와이드스크린 설정
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # 빈 슬라이드

    # 색상 정의 (Meta Brand Style)
    C_WHITE = RGBColor(255, 255, 255)
    C_BG_CARD = RGBColor(248, 250, 252)
    C_BORDER = RGBColor(226, 232, 240)
    C_TEXT_PRI = RGBColor(15, 23, 42)
    C_TEXT_SEC = RGBColor(71, 85, 105)
    C_TEXT_MUTED = RGBColor(148, 163, 184)
    C_META_BLUE = RGBColor(0, 100, 224)
    C_BLUE_LIGHT = RGBColor(235, 245, 255)
    C_EMERALD = RGBColor(5, 150, 105)
    C_EMERALD_BG = RGBColor(236, 253, 245)
    C_ROSE = RGBColor(225, 29, 72)
    C_ROSE_BG = RGBColor(255, 241, 242)
    C_CODE_BG = RGBColor(15, 23, 42)
    C_CODE_TEXT = RGBColor(226, 232, 240)

    FONT_FAMILY = "Malgun Gothic"

    def add_header(slide, tag_text, title_text, desc_text, slide_num):
        # 상단 브랜드 바
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = f"META BRAND PRESENTATION  |  SMART WAKEUP SYSTEM"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_META_BLUE

        # 페이지 번호
        p_num = tf.add_paragraph()
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"{slide_num} / 8"
        p_num.font.name = FONT_FAMILY
        p_num.font.size = Pt(10)
        p_num.font.bold = True
        p_num.font.color.rgb = C_TEXT_MUTED

        # 태그 & 타이틀
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(1.3))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        tf2.margin_top = tf2.margin_bottom = tf2.margin_left = tf2.margin_right = 0
        
        p_tag = tf2.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.name = FONT_FAMILY
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_META_BLUE

        p_title = tf2.add_paragraph()
        p_title.text = title_text
        p_title.font.name = FONT_FAMILY
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_PRI

        if desc_text:
            p_desc = tf2.add_paragraph()
            p_desc.text = desc_text
            p_desc.font.name = FONT_FAMILY
            p_desc.font.size = Pt(13)
            p_desc.font.color.rgb = C_TEXT_SEC

    # ----------------------------------------------------
    # SLIDE 1: 표지
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    # 배경 흰색 사각형
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_WHITE
    bg.line.fill.background()

    # 상단 뱃지
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_BLUE_LIGHT
    badge.line.color.rgb = C_META_BLUE
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.text = "ANTHROPIC ELI5 FRAMEWORK"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_META_BLUE
    p.alignment = PP_ALIGN.CENTER

    # 타이틀
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.7), Inches(2.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "라즈베리파이 5 담당자를 위한 첫걸음"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_PRI

    p2 = tf.add_paragraph()
    p2.text = "\"부품 없이 백엔드 통신부터 100% 끝내기!\""
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = C_META_BLUE

    p3 = tf.add_paragraph()
    p3.text = "코딩과 라즈베리파이가 처음이어도 괜찮습니다. 복잡한 전선 연결 없이 화면 출력(print)만으로 시작합니다."
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(15)
    p3.font.color.rgb = C_TEXT_SEC

    # 3개 요약 카드
    card_w = Inches(3.64)
    card_h = Inches(2.3)
    card_y = Inches(4.5)
    cards_data = [
        ("1. 하드웨어 배선 제로", "전선 쇼트나 핀 파손 걱정 NO!\n오직 파이썬 화면에 찍히는 글자(print)로 1차 성공을 확인합니다.", C_BG_CARD),
        ("2. 직관적인 우체통 비유", "5세 아이도 이해하는 식당과 우체통 비유로\nREST API(GET/POST) 통신 원리를 마스터합니다.", C_BLUE_LIGHT),
        ("3. 1분 즉석 검증", "준비된 test_connection.py 단 한 줄로\n백엔드 ↔ 파이 간 정상 통신을 즉시 증명합니다.", C_BG_CARD),
    ]
    for i, (ctitle, cdesc, cbg) in enumerate(cards_data):
        cx = Inches(0.8 + i * 4.0)
        c = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = cbg
        c.line.color.rgb = C_BORDER
        ctf = c.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.3)
        cp1 = ctf.paragraphs[0]
        cp1.text = ctitle
        cp1.font.name = FONT_FAMILY
        cp1.font.size = Pt(16)
        cp1.font.bold = True
        cp1.font.color.rgb = C_TEXT_PRI
        cp2 = ctf.add_paragraph()
        cp2.text = cdesc
        cp2.font.name = FONT_FAMILY
        cp2.font.size = Pt(12.5)
        cp2.font.color.rgb = C_TEXT_SEC

    # ----------------------------------------------------
    # SLIDE 2: 빅픽처 (식당 주방과 서빙 로봇)
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "ELI5 Big Picture", "우리는 무엇을 만드나요? (식당과 서빙 로봇의 대화)", 
               "백엔드와 라즈베리파이는 '주방 게시판'과 '서빙 로봇'의 관계와 똑같습니다.", 2)
    
    # 좌측 카드: 백엔드
    c_left = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = C_BG_CARD
    c_left.line.color.rgb = C_BORDER
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = Inches(0.35)
    
    p = tf_l.paragraphs[0]
    p.text = "🏛️ 백엔드 (중앙 우체국 & 식당 주방)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_META_BLUE

    p = tf_l.add_paragraph()
    p.text = "\n• 알람 시간 계산, 기상 미션 판정, 2차 수면 확인을 총괄하는 두뇌입니다.\n• DB(주문판)에 '지금 buzzer_1 알람 켜라!'라고 쪽지를 적어둡니다.\n• 파이가 보내온 상태 보고를 받아 프론트엔드 대시보드로 전달합니다."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.color.rgb = C_TEXT_SEC

    p = tf_l.add_paragraph()
    p.text = "\n💡 [ELI5 비유]: \"엄마가 식탁 위에 써둔 오늘의 심부름 메모장\""
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = C_META_BLUE

    # 우측 카드: 라즈베리파이
    c_right = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = C_BG_CARD
    c_right.line.color.rgb = C_BORDER
    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = Inches(0.35)

    p = tf_r.paragraphs[0]
    p.text = "🤖 라즈베리파이 5 (발 빠른 심부름꾼)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    p = tf_r.add_paragraph()
    p.text = "\n• 1초마다 우체국(백엔드)으로 달려가 새 쪽지가 있는지 확인합니다.\n• '알람 울려(ringing)' 쪽지를 발견하면 방 안에서 소리를 냅니다.\n• '터치패드가 눌렸다'는 센서 신호가 들어오면 즉시 우체국에 보고합니다."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.color.rgb = C_TEXT_SEC

    p = tf_r.add_paragraph()
    p.text = "\n💡 [ELI5 비유]: \"1초마다 식탁 쪽지를 확인하고 심부름을 해내는 착한 아이\""
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    # ----------------------------------------------------
    # SLIDE 3: 둘만의 비밀 통신 규칙 3가지
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Core Contract", "둘만의 비밀 통신 규칙 3가지 (부록 A 계약)", 
               "백엔드 친구와 라즈베리파이 담당자는 딱 이 3가지만 맞추면 서로 내부 코드를 몰라도 통신됩니다.", 3)

    rules = [
        ("1. 친구네 집 주소 (BACKEND_URL)", "http://192.168.0.xxx:8000", 
         "\"친구네 집 문패 번호예요. localhost라고 쓰면 나 자신을 가리키므로, 반드시 백엔드가 켜진 친구 PC의 실제 Wi-Fi IP를 적어야 해요!\""),
        ("2. 비밀 암호 도장 (DEVICE_API_KEY)", "wake_up_2026_09_05_v1_0_0", 
         "\"낯선 사람이 장난치지 못하도록, 쪽지 봉투 헤더(Header)에 찍는 우리 팀만의 특수 암호 스탬프예요. 철자가 1개만 틀려도 거절당해요!\""),
        ("3. 쪽지함 이름표 (device_id)", "buzzer_1 (부저) / touch_pad_1 (터치)", 
         "\"우체통 안의 칸막이 이름표예요. 백엔드 DB와 파이 코드가 똑같이 buzzer_1이라고 불러야 올바른 쪽지를 꺼내올 수 있어요!\""),
    ]

    for i, (rtitle, rval, reli5) in enumerate(rules):
        ry = Inches(2.3 + i * 1.5)
        rc = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), ry, Inches(11.7), Inches(1.3))
        rc.fill.solid()
        rc.fill.fore_color.rgb = C_BG_CARD
        rc.line.color.rgb = C_BORDER
        rtf = rc.text_frame
        rtf.word_wrap = True
        rtf.margin_left = rtf.margin_right = Inches(0.3)
        rtf.margin_top = Inches(0.2)
        
        p = rtf.paragraphs[0]
        p.text = f"{rtitle}  ➡️  {rval}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_META_BLUE

        p2 = rtf.add_paragraph()
        p2.text = reli5
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_SEC

    # ----------------------------------------------------
    # SLIDE 4: GET Polling (우체통 열어보기)
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Step 1: Polling", "우체통 열어보기 (GET desired-state)", 
               "라즈베리파이는 1초마다 백엔드에게 '저 지금 뭐 해야 하나요?'라고 물어봅니다.", 4)

    # 좌: 요청
    c_req = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_req.fill.solid()
    c_req.fill.fore_color.rgb = C_BG_CARD
    c_req.line.color.rgb = C_BORDER
    tf = c_req.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "📬 물어보는 쪽지 (GET 요청)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_META_BLUE
    p2 = tf.add_paragraph()
    p2.text = "\n[주소]\nGET /api/v1/devices/buzzer_1/desired-state\n\n[헤더 (비밀도장)]\nX-Device-Api-Key: wake_up_2026_09_05_v1_0_0"
    p2.font.name = "Consolas"
    p2.font.size = Pt(12)
    p2.font.color.rgb = C_TEXT_PRI
    p3 = tf.add_paragraph()
    p3.text = "\n💡 [ELI5]: \"우체국장님! buzzer_1 칸에 새 편지 들어온 것 있나요?\""
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(12.5)
    p3.font.bold = True
    p3.font.color.rgb = C_META_BLUE

    # 우: 응답
    c_res = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_res.fill.solid()
    c_res.fill.fore_color.rgb = C_BG_CARD
    c_res.line.color.rgb = C_BORDER
    tf = c_res.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "✉️ 돌아오는 답장 (JSON 응답: 200 OK)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p2 = tf.add_paragraph()
    p2.text = '\n{\n  "success": true,\n  "data": {\n    "device_id": "buzzer_1",\n    "desired_state": "ringing",\n    "value": {"sound": "beep"}\n  }\n}'
    p2.font.name = "Consolas"
    p2.font.size = Pt(12)
    p2.font.color.rgb = C_TEXT_PRI
    p3 = tf.add_paragraph()
    p3.text = '\n💡 [ELI5]: "응! 여기 desired_state가 ringing이라고 적혀있지? 지금 당장 시끄럽게 울려!"'
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(12.5)
    p3.font.bold = True
    p3.font.color.rgb = C_EMERALD

    # ----------------------------------------------------
    # SLIDE 5: POST Report (심부름 끝났다고 보고하기)
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Step 2: Reporting", "심부름 끝났다고 보고하기 (POST state)", 
               "알람을 켰으면, 백엔드에게 '지시하신 대로 소리 울리기 시작했습니다!'라고 알려줍니다.", 5)

    c_post = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_post.fill.solid()
    c_post.fill.fore_color.rgb = C_BG_CARD
    c_post.line.color.rgb = C_BORDER
    tf = c_post.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "📤 완료 보고서 발송 (POST 요청)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_META_BLUE
    p2 = tf.add_paragraph()
    p2.text = '\n[주소]\nPOST /api/v1/devices/buzzer_1/state\n\n[보낼 내용 (JSON)]\n{\n  "state": "ringing",\n  "reported_at": "2026-09-06T08:00:00Z"\n}'
    p2.font.name = "Consolas"
    p2.font.size = Pt(12)
    p2.font.color.rgb = C_TEXT_PRI
    p3 = tf.add_paragraph()
    p3.text = '\n💡 [ELI5]: "국장님! 방금 부저 알람 실제로 켰습니다! 완료 보고합니다!"'
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(12.5)
    p3.font.bold = True
    p3.font.color.rgb = C_META_BLUE

    c_dash = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_dash.fill.solid()
    c_dash.fill.fore_color.rgb = C_EMERALD_BG
    c_dash.line.color.rgb = RGBColor(167, 243, 208)
    tf = c_dash.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "🖥️ 웹 대시보드 화면에 실시간 반영!"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p2 = tf.add_paragraph()
    p2.text = "\n• 백엔드가 이 보고(POST)를 받는 즉시 WebSocket으로 프론트엔드에 뿌려줍니다.\n\n• 학생들의 웹 브라우저 대시보드에서 '부저 작동 중(RINGING)' 배지가 초록색으로 반짝 켜집니다!\n\n• 이로써 백엔드 ↔ 파이 ↔ 웹 대시보드 전체 고리가 완성됩니다."
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_TEXT_SEC

    # ----------------------------------------------------
    # SLIDE 6: No Hardware Philosophy (Mock의 진짜 의미)
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Phase 1 Philosophy", "빵판은 잠시 안녕! Print문으로 먼저 성공하기", 
               "전선을 꽂기 전에 화면 글자로 먼저 성공을 맛보는 것이 숙련된 엔지니어의 원칙입니다.", 6)

    c_bad = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_bad.fill.solid()
    c_bad.fill.fore_color.rgb = C_ROSE_BG
    c_bad.line.color.rgb = RGBColor(254, 205, 211)
    tf = c_bad.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "❌ 초보자가 자주 겪는 좌절과 실패"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_ROSE
    p2 = tf.add_paragraph()
    p2.text = "\n• 시작하자마자 40개짜리 GPIO 핀에 복잡하게 전선 꽂기\n• 5V와 GND를 반대로 꽂아 라즈베리파이 보드 태워먹기\n• 소리가 안 나는데, 코드가 문제인지 부품이 불량인지 몰라 멘붕\n\n⚠️ 주의: 통신이 되는지 모르는 상태에서 회로부터 연결하면 문제를 절대 못 찾습니다!"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(12.5)
    p2.font.color.rgb = C_TEXT_PRI

    c_good = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.3), Inches(5.6), Inches(4.5))
    c_good.fill.solid()
    c_good.fill.fore_color.rgb = C_EMERALD_BG
    c_good.line.color.rgb = RGBColor(167, 243, 208)
    tf = c_good.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "✅ 이번 2차 첫 단계: 100% Mock Print"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p2 = tf.add_paragraph()
    p2.text = '\n• 실제 피에조 부저 대신 print("삐-삐- 알람 울림!") 글자 출력\n• 터치패드 터치 대신 키보드 엔터 입력으로 센서 보고 시뮬레이션\n• 통신이 100% 성공하면, 그때 전선 딱 2가닥만 부저에 꽂으면 끝!\n\n🎉 목표: 터미널 화면에 성공 로그가 뜨면 오늘 라즈베리파이 임무는 100점 만점!'
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(12.5)
    p2.font.color.rgb = C_TEXT_PRI

    # ----------------------------------------------------
    # SLIDE 7: Hands-on Step (1분 만에 끝내는 첫 연결 테스트)
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Hands-on Step", "1분 만에 끝내는 첫 연결 테스트 (실습)", 
               "라즈베리파이 터미널에서 다음 딱 2줄만 입력하면 끝납니다.", 7)

    # 실행 명령어 카드
    c_cmd = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(11.7), Inches(1.3))
    c_cmd.fill.solid()
    c_cmd.fill.fore_color.rgb = C_CODE_BG
    c_cmd.line.color.rgb = C_BORDER
    tf = c_cmd.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "$ cd pi"
    p.font.name = "Consolas"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p2 = tf.add_paragraph()
    p2.text = "$ python test_connection.py"
    p2.font.name = "Consolas"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(74, 222, 128)

    # 결과 화면 카드
    c_out = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.8), Inches(11.7), Inches(3.0))
    c_out.fill.solid()
    c_out.fill.fore_color.rgb = C_BG_CARD
    c_out.line.color.rgb = C_BORDER
    tf = c_out.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "🎉 화면에 이런 글자가 나오면 대성공!"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_META_BLUE
    p2 = tf.add_paragraph()
    p2.text = """🔍 1단계: 백엔드 우체국이 문을 열었는지 확인 중...
  ✅ [성공!] 백엔드 서버와 연결되었습니다! (상태 코드: 200 OK)
📬 2단계: 1초마다 백엔드 우체통 열어보기 시작
🔔 [상태 변화 감지!] 부저 명령이 바뀌었습니다: off ➡️ ringing
  🔊 [피에조 부저 흉내내기] 삐- 삐- 삐-! 알람이 시끄럽게 울리고 있습니다!
  📤 [보고 완료] 백엔드에 '부저가 실제로 울리고 있어요'라고 보고했습니다."""
    p2.font.name = "Consolas"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = C_TEXT_PRI

    # ----------------------------------------------------
    # SLIDE 8: Blameless Troubleshooting (에러 처방전)
    # ----------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Blameless SOS", "빨간 에러가 떠도 당황 금지! 3초 처방전", 
               "에러 메시지는 내 잘못이 아니라 컴퓨터가 힌트를 주는 친절한 안내문입니다.", 8)

    errors = [
        ("Connection Refused", C_ROSE, "친구가 아직 문을 안 열어줬어요.", 
         "1. 백엔드 PC에서 서버가 켜져 있나요?\n2. 파이와 백엔드 PC가 같은 Wi-Fi에 연결되어 있나요?\n3. pi/.env 파일의 IP 주소를 다시 확인하세요."),
        ("404 Not Found", RGBColor(217, 119, 6), "우체통에 적힌 이름표가 달라요.", 
         "1. 백엔드 DB의 devices 테이블에 'buzzer_1'이 등록되어 있는지 백엔드 친구에게 물어보세요!\n2. 대소문자나 철자를 비교해보세요."),
        ("401 / 403 Forbidden", RGBColor(99, 102, 241), "비밀 암호 도장이 틀렸어요.", 
         "1. pi/.env의 DEVICE_API_KEY와 백엔드 .env의 DEVICE_API_KEY가 완전히 똑같은지 비교하세요!\n(예: wake_up_2026_09_05_v1_0_0)"),
    ]

    for i, (err_name, err_col, err_eli5, err_fix) in enumerate(errors):
        ex = Inches(0.8 + i * 4.0)
        ec = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ex, Inches(2.3), card_w, Inches(4.5))
        ec.fill.solid()
        ec.fill.fore_color.rgb = C_BG_CARD
        ec.line.color.rgb = C_BORDER
        etf = ec.text_frame
        etf.word_wrap = True
        etf.margin_left = etf.margin_right = Inches(0.25)
        etf.margin_top = Inches(0.3)

        p = etf.paragraphs[0]
        p.text = err_name
        p.font.name = FONT_FAMILY
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = err_col

        p2 = etf.add_paragraph()
        p2.text = f"\n💡 [비유]: {err_eli5}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_PRI

        p3 = etf.add_paragraph()
        p3.text = f"\n[3초 처방전]:\n{err_fix}"
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(11.5)
        p3.font.color.rgb = C_TEXT_SEC

    # 저장
    output_path = "docs/라즈베리파이5_백엔드_연동_ELI5_가이드.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
